# -*- coding: utf-8 -*-
"""Genera Opciones_Ejercicio7.xlsx: una hoja por estrategia (formato Ejercicio 7,
con formulas vivas) + hoja resumen del portafolio de opciones."""
import numpy as np
import data as d
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

R = 0.045  # tasa libre de riesgo anual (T-bill ~3m, ilustrativa)

# ---- estilos ----
HDR = Font(bold=True, color="FFFFFF", size=11)
HDRFILL = PatternFill("solid", fgColor="1F4E78")
SUB = Font(bold=True, color="1F4E78")
TITLE = Font(bold=True, size=14, color="1F4E78")
CENTER = Alignment(horizontal="center", vertical="center")
BORDER = Border(*[Side(style="thin", color="BFBFBF")] * 4)
GREEN = PatternFill("solid", fgColor="C6EFCE")
RED = PatternFill("solid", fgColor="FFC7CE")
YEL = PatternFill("solid", fgColor="FFF2CC")
GREY = PatternFill("solid", fgColor="D9D9D9")

def style_header_row(ws, row, c1, c2):
    for c in range(c1, c2 + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HDR; cell.fill = HDRFILL; cell.alignment = CENTER
        cell.border = BORDER

def T_years(exp):
    from datetime import date
    y, m, dd = map(int, exp.split("-"))
    return max((date(y, m, dd) - date(2026, 7, 2)).days, 0) / 365.0

wb = Workbook()
wb.remove(wb.active)

# ============================================================ hojas por estrategia
for pos in d.OPCIONES:
    tick = pos["subyacente"]
    ws = wb.create_sheet(pos["id"][:31])
    ws.sheet_view.showGridLines = False
    spot = d.SPOT[tick]
    legs = pos["legs"]
    Tv = T_years(pos["exp"])

    ws["A1"] = f"{tick} — {pos['nombre']}"
    ws["A1"].font = TITLE
    ws["A2"] = f"{pos['estrategia']}  ·  Sesgo: {pos['sesgo']}  ·  Confianza: {pos['confianza']}"
    ws["A2"].font = Font(italic=True, color="595959")

    # ---- bloque de parametros ----
    ws["A4"] = "PARAMETROS"; ws["A4"].font = SUB
    params = [
        ("Subyacente", tick),
        ("Spot de entrada (2-jul-2026)", spot),
        ("Contratos", pos["contratos"]),
        ("Acciones por contrato", 100),
        ("Total de acciones", "=B7*B8"),
        ("Vencimiento", pos["exp"]),
        ("r (anual)", R),
        ("T (anios)", round(Tv, 4)),
    ]
    r0 = 5
    for i, (lab, val) in enumerate(params):
        ws.cell(row=r0 + i, column=1, value=lab).font = Font(bold=True)
        ws.cell(row=r0 + i, column=2, value=val)
    # celdas clave: B6=spot, B7=contratos, B8=100
    B_CONTR, B_MULT = "B7", "B8"

    # ---- tabla de patas ----
    rleg_hdr = r0 + len(params) + 1  # header row for legs
    ws.cell(row=rleg_hdr, column=1, value="PATAS DE LA ESTRATEGIA").font = SUB
    hrow = rleg_hdr + 1
    leg_cols = ["Pata", "Tipo", "Lado", "Strike (K)", "Prima/accion", "Mult"]
    for j, h in enumerate(leg_cols):
        ws.cell(row=hrow, column=1 + j, value=h)
    style_header_row(ws, hrow, 1, len(leg_cols))
    leg_rows = []
    for k, lg in enumerate(legs):
        rr = hrow + 1 + k
        leg_rows.append(rr)
        ws.cell(row=rr, column=1, value=f"P{k+1}")
        ws.cell(row=rr, column=2, value=lg["tipo"].upper())
        ws.cell(row=rr, column=3, value=lg["lado"].upper())
        ws.cell(row=rr, column=4, value=lg["K"])
        ws.cell(row=rr, column=5, value=lg["prima"])
        ws.cell(row=rr, column=6, value=lg.get("mult", 1))
        for c in range(1, 7):
            ws.cell(row=rr, column=c).border = BORDER

    # prima neta (credito + / debito -) por accion y total, con formula viva
    # sum: short -> +prima*mult ; long -> -prima*mult
    terms = []
    for rr, lg in zip(leg_rows, legs):
        sgn = "+" if lg["lado"] == "short" else "-"
        terms.append(f"{sgn}E{rr}*F{rr}")
    rneto = leg_rows[-1] + 2
    ws.cell(row=rneto, column=1, value="Prima neta / accion (cred + / deb -)").font = Font(bold=True)
    ws.cell(row=rneto, column=5, value="=" + "".join(terms).lstrip("+"))
    ws.cell(row=rneto + 1, column=1, value="Prima neta total (USD)").font = Font(bold=True)
    ws.cell(row=rneto + 1, column=5, value=f"=E{rneto}*{B_CONTR}*{B_MULT}")
    ws.cell(row=rneto, column=5).number_format = "0.00"
    ws.cell(row=rneto + 1, column=5).number_format = "#,##0"

    # ---- tabla de payoff (formato Ejercicio 7) ----
    tstart = rneto + 3
    ws.cell(row=tstart, column=1, value="TABLA DE PAYOFF AL VENCIMIENTO").font = SUB
    hr = tstart + 1
    # columnas: ST | por pata: K, Prima, Moneyness, Ejerce, Payoff/acc, GoP/acc | Payoff neto/acc | GoP neto/acc | GoP TOTAL
    headers = ["ST (Spot)"]
    per_leg = ["K", "Prima", "Moneyness", "Ejerce", "Payoff/acc", "GoP/acc"]
    for k in range(len(legs)):
        headers += [f"P{k+1} {h}" for h in per_leg]
    headers += ["Payoff neto/acc", "GoP neto/acc", "GoP TOTAL (USD)"]
    for j, h in enumerate(headers):
        ws.cell(row=hr, column=1 + j, value=h)
    style_header_row(ws, hr, 1, len(headers))

    # rango de ST alrededor de strikes
    Ks = [lg["K"] for lg in legs]
    lo, hi = min(Ks), max(Ks)
    span = max(hi - lo, spot * 0.06)
    ST_vals = list(np.round(np.linspace(lo - span * 1.2 - 2, hi + span * 1.2 + 2, 23), 2))
    # asegurar que strikes y spot esten
    for extra in Ks + [round(spot, 2)]:
        if extra not in ST_vals:
            ST_vals.append(extra)
    ST_vals = sorted(set(ST_vals))

    gop_total_first = hr + 1
    for i, ST in enumerate(ST_vals):
        r = hr + 1 + i
        ws.cell(row=r, column=1, value=ST).number_format = "0.00"
        col = 2
        payoff_cells = []
        gop_cells = []
        for rr, lg in zip(leg_rows, legs):
            Kc, Pc, Fc = f"$D${rr}", f"$E${rr}", f"$F${rr}"
            STc = f"$A{r}"
            # K, Prima (referencias)
            ws.cell(row=r, column=col, value=f"={Kc}"); ws.cell(row=r, column=col).number_format = "0.00"
            ws.cell(row=r, column=col + 1, value=f"={Pc}"); ws.cell(row=r, column=col + 1).number_format = "0.00"
            if lg["tipo"] == "call":
                mny = f'=IF({STc}>{Kc},"ITM",IF({STc}={Kc},"ATM","OTM"))'
                ejc = f'=IF({STc}>{Kc},"SI","NO")'
                base = f"MAX({STc}-{Kc},0)"
            else:
                mny = f'=IF({STc}<{Kc},"ITM",IF({STc}={Kc},"ATM","OTM"))'
                ejc = f'=IF({STc}<{Kc},"SI","NO")'
                base = f"MAX({Kc}-{STc},0)"
            sgn = "" if lg["lado"] == "long" else "-"
            payoff_f = f"={sgn}{base}*{Fc}"
            pay_cell = f"{get_column_letter(col+4)}{r}"
            if lg["lado"] == "long":
                gop_f = f"={pay_cell}-{Pc}*{Fc}"
            else:
                gop_f = f"={Pc}*{Fc}+{pay_cell}"
            ws.cell(row=r, column=col + 2, value=mny)
            ws.cell(row=r, column=col + 3, value=ejc)
            ws.cell(row=r, column=col + 4, value=payoff_f).number_format = "0.00"
            ws.cell(row=r, column=col + 5, value=gop_f).number_format = "0.00"
            payoff_cells.append(pay_cell)
            gop_cells.append(f"{get_column_letter(col+5)}{r}")
            col += 6
        # neto
        ws.cell(row=r, column=col, value="=" + "+".join(payoff_cells)).number_format = "0.00"
        neto_gop_cell = f"{get_column_letter(col+1)}{r}"
        ws.cell(row=r, column=col + 1, value="=" + "+".join(gop_cells)).number_format = "0.00"
        ws.cell(row=r, column=col + 2, value=f"={neto_gop_cell}*{B_CONTR}*{B_MULT}").number_format = "#,##0"
        # color GoP total
        gp_val = d.estrategia_gp_total(pos, ST)
        ws.cell(row=r, column=col + 2).fill = GREEN if gp_val >= 0 else RED
        for c in range(1, len(headers) + 1):
            ws.cell(row=r, column=c).border = BORDER
    gop_total_last = hr + len(ST_vals)
    gcol = get_column_letter(len(headers))

    # ---- resumen: BE, maxG, maxL ----
    rs = gop_total_last + 2
    ws.cell(row=rs, column=1, value="RESUMEN DE LA ESTRATEGIA").font = SUB
    # break-even calc
    STdense = np.linspace(min(ST_vals) - 5, max(ST_vals) + 5, 4000)
    gps = np.array([d.estrategia_gp_total(pos, s) for s in STdense])
    bes = []
    for i in range(1, len(STdense)):
        if gps[i - 1] == 0 or (gps[i - 1] < 0) != (gps[i] < 0):
            bes.append(round(STdense[i], 2))
    bes = sorted(set(round(b, 2) for b in bes))
    resumen = [
        ("Break-even(s)", ", ".join(f"{b:.2f}" for b in bes) if bes else "n/a"),
        ("Ganancia maxima (USD)", f"=MAX({gcol}{gop_total_first}:{gcol}{gop_total_last})"),
        ("Perdida maxima (USD)", f"=MIN({gcol}{gop_total_first}:{gcol}{gop_total_last})"),
        ("Prima neta total (USD)", f"=E{rneto+1}"),
        ("Capital en riesgo (USD)", round(-min(d.estrategia_gp_total(pos, s) for s in STdense), 0)),
    ]
    for i, (lab, val) in enumerate(resumen):
        ws.cell(row=rs + 1 + i, column=1, value=lab).font = Font(bold=True)
        c = ws.cell(row=rs + 1 + i, column=3, value=val)
        if isinstance(val, str) and val.startswith("="):
            c.number_format = "#,##0"

    # ---- escenarios optimista/base/pesimista (valores) ----
    re = rs + len(resumen) + 2
    ws.cell(row=re, column=1, value="RESULTADO POR ESCENARIO (al vencimiento)").font = SUB
    esc_hdr = re + 1
    for j, h in enumerate(["Escenario", "ST asumido", "GoP total (USD)", "Rent. s/ capital riesgo"]):
        ws.cell(row=esc_hdr, column=1 + j, value=h)
    style_header_row(ws, esc_hdr, 1, 4)
    cap_riesgo = -min(d.estrategia_gp_total(pos, s) for s in STdense) or 1
    if pos["sesgo"].startswith("Bajista"):
        escs = [("Optimista", spot * 0.94), ("Base", spot), ("Pesimista", spot * 1.06)]
    elif pos["sesgo"].startswith("Neutral"):
        mid = np.mean(Ks)
        escs = [("Optimista", mid), ("Base", spot), ("Pesimista", spot * 1.06)]
    else:
        escs = [("Optimista", spot * 1.06), ("Base", spot), ("Pesimista", spot * 0.94)]
    for j, (nm, st) in enumerate(escs):
        gp = d.estrategia_gp_total(pos, st)
        rr = esc_hdr + 1 + j
        ws.cell(row=rr, column=1, value=nm)
        ws.cell(row=rr, column=2, value=round(st, 2)).number_format = "0.00"
        cc = ws.cell(row=rr, column=3, value=round(gp, 0)); cc.number_format = "#,##0"
        cc.fill = GREEN if gp >= 0 else RED
        ws.cell(row=rr, column=4, value=gp / cap_riesgo).number_format = "0.0%"
        for c in range(1, 5):
            ws.cell(row=rr, column=c).border = BORDER

    # anchos
    ws.column_dimensions["A"].width = 30
    for c in range(2, len(headers) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 12

# ============================================================ hoja resumen
ws = wb.create_sheet("RESUMEN_OPCIONES", 0)
ws.sheet_view.showGridLines = False
ws["A1"] = "PORTAFOLIO DE OPCIONES — Resumen (entrada 2-jul-2026)"
ws["A1"].font = TITLE
cols = ["Activo", "Derivado", "Estrategia", "Posicion/Sesgo", "Spot", "Strikes",
        "Prima neta/acc", "Contratos", "Valor nocional", "Cap. en riesgo", "Peso riesgo"]
hr = 3
for j, h in enumerate(cols):
    ws.cell(row=hr, column=1 + j, value=h)
style_header_row(ws, hr, 1, len(cols))

STdense_cache = {}
riesgos = []
for pos in d.OPCIONES:
    STs = np.linspace(min(l["K"] for l in pos["legs"]) - 40,
                      max(l["K"] for l in pos["legs"]) + 40, 5000)
    cr = -min(d.estrategia_gp_total(pos, s) for s in STs)
    riesgos.append(cr)
tot_riesgo = sum(riesgos)

r = hr + 1
for pos, cr in zip(d.OPCIONES, riesgos):
    tick = pos["subyacente"]
    spot = d.SPOT[tick]
    neto = d.credito_debito_neto(pos)
    nocional = spot * 100 * pos["contratos"]
    strikes = "/".join(str(int(l["K"]) if l["K"] == int(l["K"]) else l["K"]) for l in pos["legs"])
    vals = ["Opcion" if False else tick, "Opcion", pos["estrategia"], pos["sesgo"], spot,
            strikes, round(neto, 2), pos["contratos"], round(nocional, 0),
            round(cr, 0), cr / tot_riesgo]
    vals[0] = tick
    for j, v in enumerate(vals):
        c = ws.cell(row=r, column=1 + j, value=v)
        c.border = BORDER
    ws.cell(row=r, column=9).number_format = "#,##0"
    ws.cell(row=r, column=10).number_format = "#,##0"
    ws.cell(row=r, column=7).number_format = "0.00"
    ws.cell(row=r, column=11).number_format = "0.0%"
    r += 1
# totales
ws.cell(row=r, column=1, value="TOTAL").font = Font(bold=True)
ws.cell(row=r, column=9, value=f"=SUM(I{hr+1}:I{r-1})").number_format = "#,##0"
ws.cell(row=r, column=10, value=f"=SUM(J{hr+1}:J{r-1})").number_format = "#,##0"
ws.cell(row=r, column=11, value=f"=SUM(K{hr+1}:K{r-1})").number_format = "0.0%"
for c in range(1, len(cols) + 1):
    ws.cell(row=r, column=c).fill = GREY

# distribuciones
r2 = r + 3
ws.cell(row=r2, column=1, value="DISTRIBUCION POR ESTRATEGIA (por capital en riesgo)").font = SUB
from collections import defaultdict
dist_estrat = defaultdict(float); dist_sector = defaultdict(float); dist_riesgo = defaultdict(float)
for pos, cr in zip(d.OPCIONES, riesgos):
    dist_estrat[pos["estrategia"]] += cr
    dist_sector[pos["sector"]] += cr
    dist_riesgo[pos["riesgo_nivel"]] += cr

def dump_dist(ws, start, title, dist):
    ws.cell(row=start, column=1, value=title).font = SUB
    ws.cell(row=start + 1, column=1, value="Categoria").font = HDR
    ws.cell(row=start + 1, column=2, value="Cap. riesgo").font = HDR
    ws.cell(row=start + 1, column=3, value="Peso").font = HDR
    style_header_row(ws, start + 1, 1, 3)
    rr = start + 2
    for k, v in sorted(dist.items(), key=lambda x: -x[1]):
        ws.cell(row=rr, column=1, value=k).border = BORDER
        ws.cell(row=rr, column=2, value=round(v, 0)).number_format = "#,##0"
        ws.cell(row=rr, column=2).border = BORDER
        ws.cell(row=rr, column=3, value=v / tot_riesgo).number_format = "0.0%"
        ws.cell(row=rr, column=3).border = BORDER
        rr += 1
    return rr + 2

nxt = dump_dist(ws, r2 + 1, "DISTRIBUCION POR ESTRATEGIA", dist_estrat)
nxt = dump_dist(ws, nxt, "DISTRIBUCION POR SECTOR", dist_sector)
nxt = dump_dist(ws, nxt, "DISTRIBUCION POR NIVEL DE RIESGO", dist_riesgo)

ws.column_dimensions["A"].width = 34
for c in range(2, len(cols) + 1):
    ws.column_dimensions[get_column_letter(c)].width = 15

wb.save("Opciones_Ejercicio7.xlsx")
print("OK Opciones_Ejercicio7.xlsx  hojas:", wb.sheetnames)
