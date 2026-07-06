# -*- coding: utf-8 -*-
"""Genera Futuros_MarcaAMercado.xlsx: una hoja por futuro (mark-to-market con
cuenta de margen y margin call, formulas vivas) + hoja resumen."""
import data as d
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HDR = Font(bold=True, color="FFFFFF", size=11)
HDRFILL = PatternFill("solid", fgColor="1F4E78")
SUB = Font(bold=True, color="1F4E78")
TITLE = Font(bold=True, size=14, color="1F4E78")
CENTER = Alignment(horizontal="center", vertical="center")
BORDER = Border(*[Side(style="thin", color="BFBFBF")] * 4)
GREEN = PatternFill("solid", fgColor="C6EFCE")
RED = PatternFill("solid", fgColor="FFC7CE")
CALLFILL = PatternFill("solid", fgColor="FFC000")
GREY = PatternFill("solid", fgColor="D9D9D9")

def style_header_row(ws, row, c1, c2):
    for c in range(c1, c2 + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HDR; cell.fill = HDRFILL; cell.alignment = CENTER; cell.border = BORDER

wb = Workbook(); wb.remove(wb.active)

resumen_rows = []

for sym, f in d.FUTUROS.items():
    ws = wb.create_sheet(sym)
    ws.sheet_view.showGridLines = False
    ws["A1"] = f"{f['activo']} ({sym}) — {f['mercado']}  ·  Marca a Mercado"
    ws["A1"].font = TITLE
    ws["A2"] = f"Posicion: {f['posicion']}  ·  Tesis: {f['tesis']}"
    ws["A2"].font = Font(italic=True, color="595959")

    # ---- parametros ----
    ws["A4"] = "PARAMETROS"; ws["A4"].font = SUB
    sign = 1 if f["posicion"] == "Larga" else -1
    P = [
        ("Activo", f["activo"]), ("Mercado", f["mercado"]),
        ("Posicion (Larga=1 / Corta=-1)", sign),
        ("N.o de contratos", f["contratos"]),
        ("Tamano de contrato", f["tamano"]),
        ("Tamano del portafolio (unidades)", "=B8*B9"),
        ("Margen inicial (por contrato)", f["margen_ini"]),
        ("Margen inicial (total)", "=B8*B11"),
        ("Margen de mantenimiento (por contrato)", f["margen_mant"]),
        ("Margen de mantenimiento (total)", "=B8*B13"),
        ("Precio negociado de entrada", f["entrada"]),
    ]
    for i, (lab, val) in enumerate(P):
        ws.cell(row=5 + i, column=1, value=lab).font = Font(bold=True)
        c = ws.cell(row=5 + i, column=2, value=val)
        if isinstance(val, (int, float)) and abs(val) > 100:
            c.number_format = "#,##0.00"
    # celdas clave
    SIGN, NC, TC, UNID = "$B$7_", "$B$7", "$B$8", "$B$9"
    C_SIGN = "$B$7"  # careful: recompute below
    # actual mapping:
    # B7 sign, B8 contratos, B9 tamano, B10 unidades, B11 marg ini/contr, B12 marg ini tot,
    # B13 marg mant/contr, B14 marg mant tot, B15 entrada
    C_SIGN, C_NC, C_TC, C_UNID = "$B$7", "$B$8", "$B$9", "$B$10"
    C_INI_TOT, C_MANT_TOT, C_ENTRADA = "$B$12", "$B$14", "$B$15"

    # ---- tabla diaria ----
    tstart = 18
    ws.cell(row=tstart, column=1, value="MARCA A MERCADO DIARIA").font = SUB
    ws.cell(row=tstart, column=6,
            value="Precios de liquidacion ILUSTRATIVOS (reemplazar por settlement oficial "
                  + f["mercado"] + ")").font = Font(italic=True, color="C00000", size=9)
    hr = tstart + 1
    cols = ["Fecha", "Dia", "Precio negociado", "Precio liquidacion",
            "G/P por unidad", "G/P total diaria", "G/P acumulada",
            "Cuenta de Margen", "Margin Call"]
    for j, h in enumerate(cols):
        ws.cell(row=hr, column=1 + j, value=h)
    style_header_row(ws, hr, 1, len(cols))

    first = hr + 1
    for i, (fecha, dia) in enumerate(d.FECHAS_FUT):
        r = first + i
        settle = f["settle"][i]
        ws.cell(row=r, column=1, value=fecha)
        ws.cell(row=r, column=2, value=dia)
        # precio negociado: entrada solo el primer dia
        if i == 0:
            ws.cell(row=r, column=3, value=f"={C_ENTRADA}").number_format = "#,##0.00"
            prev = C_ENTRADA
        else:
            ws.cell(row=r, column=3, value="—")
            prev = f"$D${r-1}"
        ws.cell(row=r, column=4, value=settle).number_format = "#,##0.00"
        # G/P por unidad = (settle - prev)*sign
        ws.cell(row=r, column=5, value=f"=(D{r}-{prev})*{C_SIGN}").number_format = "#,##0.00"
        # G/P total diaria = E * unidades
        ws.cell(row=r, column=6, value=f"=E{r}*{C_UNID}").number_format = "#,##0"
        # G/P acumulada
        if i == 0:
            ws.cell(row=r, column=7, value=f"=F{r}").number_format = "#,##0"
        else:
            ws.cell(row=r, column=7, value=f"=G{r-1}+F{r}").number_format = "#,##0"
        # Cuenta de margen (con top-up) y Margin Call
        prev_acct = C_INI_TOT if i == 0 else f"$H${r-1}"
        # margin call = IF(prev_acct + F < mant, ini - (prev_acct+F), 0)
        ws.cell(row=r, column=9,
                value=f"=IF({prev_acct}+F{r}<{C_MANT_TOT},{C_INI_TOT}-({prev_acct}+F{r}),0)"
                ).number_format = "#,##0"
        # cuenta = prev_acct + F + call
        ws.cell(row=r, column=8, value=f"={prev_acct}+F{r}+I{r}").number_format = "#,##0"
        # formato condicional manual (color por signo, resaltar call)
        gp_dia = (settle - (f["entrada"] if i == 0 else f["settle"][i-1])) * sign * f["contratos"] * f["tamano"]
        ws.cell(row=r, column=6).fill = GREEN if gp_dia >= 0 else RED
        # margin call highlight
        # recompute account to know if call fired
        for c in range(1, len(cols) + 1):
            ws.cell(row=r, column=c).border = BORDER
    last = first + len(d.FECHAS_FUT) - 1

    # resaltar celdas de margin call > 0
    from openpyxl import load_workbook  # noqa
    # (marcamos por valor recalculado en python)
    cuenta = f["contratos"] * f["margen_ini"]; prev_s = f["entrada"]
    ini_tot = f["contratos"] * f["margen_ini"]; mant_tot = f["contratos"] * f["margen_mant"]
    for i, s in enumerate(f["settle"]):
        r = first + i
        gpt = (s - prev_s) * sign * f["contratos"] * f["tamano"]
        antes = cuenta + gpt
        call = (ini_tot - antes) if antes < mant_tot else 0
        if call > 0:
            ws.cell(row=r, column=9).fill = CALLFILL
            ws.cell(row=r, column=9).font = Font(bold=True, color="C00000")
        cuenta = antes + call; prev_s = s

    # ---- resumen del futuro ----
    rs = last + 2
    ws.cell(row=rs, column=1, value="RESUMEN").font = SUB
    res = [
        ("G/P acumulada total (USD)", f"=G{last}"),
        ("Reposiciones por Margin Call (USD)", f"=SUM(I{first}:I{last})"),
        ("Cuenta de margen final (USD)", f"=H{last}"),
        ("Valor nocional inicial (USD)", f"={C_ENTRADA}*{C_UNID}"),
        ("Margen inicial total (USD)", f"={C_INI_TOT}"),
        ("Retorno s/ margen inicial", f"=G{last}/{C_INI_TOT}"),
    ]
    for i, (lab, val) in enumerate(res):
        ws.cell(row=rs + 1 + i, column=1, value=lab).font = Font(bold=True)
        c = ws.cell(row=rs + 1 + i, column=3, value=val)
        c.number_format = "0.0%" if "Retorno" in lab else "#,##0"

    ws.column_dimensions["A"].width = 34
    for c in range(2, len(cols) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 16

    # datos para resumen global (python)
    resumen_rows.append({
        "sym": sym, "activo": f["activo"], "mercado": f["mercado"],
        "pos": f["posicion"], "contr": f["contratos"], "tamano": f["tamano"],
        "nocional": f["entrada"] * f["contratos"] * f["tamano"],
        "ini": ini_tot, "pl": None,  # se calcula abajo
    })

# recomputar P&L en python para resumen
for row, (sym, f) in zip(resumen_rows, d.FUTUROS.items()):
    sign = 1 if f["posicion"] == "Larga" else -1
    prev = f["entrada"]; acum = 0
    for s in f["settle"]:
        acum += (s - prev) * sign * f["contratos"] * f["tamano"]; prev = s
    row["pl"] = acum

# ============================================================ hoja resumen
ws = wb.create_sheet("RESUMEN_FUTUROS", 0)
ws.sheet_view.showGridLines = False
ws["A1"] = "PORTAFOLIO DE FUTUROS — Resumen (18-jun a 9-jul-2026)"
ws["A1"].font = TITLE
ws["A2"] = "Marca a mercado con cuenta de margen. Settlements ilustrativos (ver nota metodologica)."
ws["A2"].font = Font(italic=True, color="C00000")
cols = ["Activo", "Simbolo", "Mercado", "Posicion", "Contratos", "Tamano",
        "Valor nocional", "Margen inicial", "G/P acumulada", "Retorno s/margen", "Peso nocional"]
hr = 4
for j, h in enumerate(cols):
    ws.cell(row=hr, column=1 + j, value=h)
style_header_row(ws, hr, 1, len(cols))
tot_noc = sum(r["nocional"] for r in resumen_rows)
r = hr + 1
for row in resumen_rows:
    vals = [row["activo"], row["sym"], row["mercado"], row["pos"], row["contr"],
            row["tamano"], round(row["nocional"], 0), round(row["ini"], 0),
            round(row["pl"], 0), row["pl"] / row["ini"], row["nocional"] / tot_noc]
    for j, v in enumerate(vals):
        c = ws.cell(row=r, column=1 + j, value=v); c.border = BORDER
    ws.cell(row=r, column=7).number_format = "#,##0"
    ws.cell(row=r, column=8).number_format = "#,##0"
    c = ws.cell(row=r, column=9); c.number_format = "#,##0"; c.fill = GREEN if row["pl"] >= 0 else RED
    ws.cell(row=r, column=10).number_format = "0.0%"
    ws.cell(row=r, column=11).number_format = "0.0%"
    r += 1
ws.cell(row=r, column=1, value="TOTAL").font = Font(bold=True)
ws.cell(row=r, column=7, value=f"=SUM(G{hr+1}:G{r-1})").number_format = "#,##0"
ws.cell(row=r, column=8, value=f"=SUM(H{hr+1}:H{r-1})").number_format = "#,##0"
ws.cell(row=r, column=9, value=f"=SUM(I{hr+1}:I{r-1})").number_format = "#,##0"
ws.cell(row=r, column=11, value=f"=SUM(K{hr+1}:K{r-1})").number_format = "0.0%"
for c in range(1, len(cols) + 1):
    ws.cell(row=r, column=c).fill = GREY
ws.column_dimensions["A"].width = 20
for c in range(2, len(cols) + 1):
    ws.column_dimensions[get_column_letter(c)].width = 15

wb.save("Futuros_MarcaAMercado.xlsx")
print("OK Futuros_MarcaAMercado.xlsx  hojas:", wb.sheetnames)
