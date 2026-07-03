#!/usr/bin/env python3
"""Genera CapituloII_Fretel_Anderson_v2.docx a partir del markdown fuente,
con formato de tesis (Times New Roman 12, interlineado 1.5, texto justificado)."""
import re
from pathlib import Path

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

BASE = Path(__file__).parent
SRC = BASE / "CapituloII_Fretel_Anderson_v2.md"
OUT = BASE / "CapituloII_Fretel_Anderson_v2.docx"

FONT = "Times New Roman"


def set_style(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(6)
    for name, size in [("Heading 1", 14), ("Heading 2", 13), ("Heading 3", 12), ("Heading 4", 12)]:
        st = doc.styles[name]
        st.font.name = FONT
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = None
        rpr = st.element.get_or_add_rPr()
        rpr_fonts = rpr.get_or_add_rFonts()
        rpr_fonts.set(qn("w:ascii"), FONT)
        rpr_fonts.set(qn("w:hAnsi"), FONT)


def add_runs(par, text):
    """Interpreta **negrita** y *cursiva* dentro de un párrafo."""
    tokens = re.split(r"(\*\*.+?\*\*|\*[^*]+?\*)", text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            run = par.add_run(tok[2:-2])
            run.bold = True
        elif tok.startswith("*") and tok.endswith("*"):
            run = par.add_run(tok[1:-1])
            run.italic = True
        else:
            par.add_run(tok)


def cover(doc):
    lines = [
        ("Universidad Nacional Mayor de San Marcos", 14, True),
        ("Facultad de Ciencias Económicas – Escuela Profesional de Economía", 12, False),
        ("Proyecto de Tesis de Grado", 12, False),
        ("", 12, False),
        ("CAPÍTULO II", 16, True),
        ("MARCO TEÓRICO E HIPÓTESIS", 14, True),
        ("", 12, False),
        ('"El efecto de las macrorregiones sobre los ingresos laborales en el Perú: '
         'evidencia de la ENAHO, 2021-2023"', 12, False),
        ("", 12, False),
        ("Autor: Anderson Riquelme Fretel Rojas", 12, False),
        ("Docente: Oscar Chavez Polo", 12, False),
        ("Lima, Perú – 2026", 12, False),
    ]
    for text, size, bold in lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.font.size = Pt(size)
        run.bold = bold
    doc.add_page_break()


def build_table(doc, rows):
    header = [c.strip() for c in rows[0].strip("|").split("|")]
    body = [[c.strip() for c in r.strip("|").split("|")] for r in rows[2:]]
    table = doc.add_table(rows=1 + len(body), cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, text in enumerate(header):
        cell = table.rows[0].cells[j]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(re.sub(r"\*\*", "", text))
        run.bold = True
        run.font.size = Pt(10)
    for i, row in enumerate(body, start=1):
        section_row = row[0].startswith("**") and all(not c for c in row[1:])
        for j, text in enumerate(row):
            if j >= len(header):
                break
            cell = table.rows[i].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            clean = re.sub(r"\*\*", "", text)
            run = p.add_run(clean)
            run.font.size = Pt(10)
            if section_row:
                run.bold = True
    doc.add_paragraph()


def main():
    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(3)
        section.right_margin = Cm(2.5)
    set_style(doc)
    cover(doc)

    lines = SRC.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i].strip())
                i += 1
            build_table(doc, rows)
            continue
        if line.startswith("#### "):
            doc.add_heading(line[5:], level=4)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("$$") and line.endswith("$$"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(line[2:-2].strip())
            run.italic = True
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_runs(p, line[2:])
        else:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_runs(p, line)
        i += 1

    doc.save(OUT)
    print(f"OK -> {OUT}")


if __name__ == "__main__":
    main()
