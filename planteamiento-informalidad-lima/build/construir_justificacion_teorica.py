# -*- coding: utf-8 -*-
"""1.2.1 Justificación teórica (~2 páginas) en Word (.docx).

Incluye el texto del autor (en negro) y los aportes de Claude (en ROJO), con sus
citas y referencias en APA 7. Tres teorías añadidas:
  1) El nivel de ingresos contribuye a disminuir la informalidad (enfoque costo-beneficio).
  2) A mayor nivel de ingresos, menor informalidad (enfoque estructural/desarrollo).
  3) Género, edad y estado civil influyen en la informalidad (capital humano, discriminación
     y economía de la familia).

Formato: A4, Times New Roman 12, interlineado 1.5, justificado, márgenes 2.54 cm.
"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn

BASE = os.path.dirname(__file__)
OUT = os.path.join(os.path.dirname(BASE), "1.2.1_Justificacion_Teorica.docx")

AZUL = RGBColor(0x1F, 0x4E, 0x79)
ROJO = RGBColor(0xC0, 0x00, 0x00)
NEGRO = RGBColor(0x00, 0x00, 0x00)

doc = Document()

normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(12)
rpr = normal.element.get_or_add_rPr()
rf = rpr.get_or_add_rFonts()
rf.set(qn("w:ascii"), "Times New Roman")
rf.set(qn("w:hAnsi"), "Times New Roman")
rf.set(qn("w:cs"), "Times New Roman")
pf = normal.paragraph_format
pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
pf.line_spacing = 1.5
pf.space_after = Pt(0)
pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

for sec in doc.sections:
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(2.54)
    sec.right_margin = Cm(2.54)


def _font(run, name="Times New Roman"):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rf = rpr.get_or_add_rFonts()
    rf.set(qn("w:ascii"), name)
    rf.set(qn("w:hAnsi"), name)
    rf.set(qn("w:cs"), name)


def heading(text, size=13):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(size)
    r.font.color.rgb = AZUL
    _font(r)
    return p


def para(text, color=NEGRO, after=4, indent=True):
    """Párrafo justificado con color. **negrita** soportada."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(after)
    if indent:
        p.paragraph_format.first_line_indent = Cm(1.25)
    for i, seg in enumerate(text.split("**")):
        r = p.add_run(seg)
        _font(r)
        r.font.color.rgb = color
        if i % 2 == 1:
            r.bold = True
    return p


def referencia(text, color=NEGRO):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Cm(1.25)
    p.paragraph_format.first_line_indent = Cm(-1.25)
    for i, seg in enumerate(text.split("//")):
        r = p.add_run(seg)
        _font(r)
        r.font.color.rgb = color
        if i % 2 == 1:
            r.italic = True


# ======================================================================
heading("1.2.1. Justificación teórica", size=14)

# ---------------- TEXTO DEL AUTOR (negro) ----------------
para(
    "De acuerdo a la teoría dualista, planteada por Hart (1973) y Tokman (1978), señalan "
    "que el sector no formal involucra actividades marginales que difieren "
    "significativamente de las que se encuentran en el sector formal de la economía.")

para(
    "De acuerdo con Uribe et al. (2006), la informalidad laboral se percibe como el "
    "efecto que se obtiene debido a que existe un limitado desarrollo en el sector "
    "económico moderno, con lo que no es posible para este incorporar a todo el mercado "
    "laboral que se encuentra en ese momento.")

para(
    "La teoría de la segmentación del mercado laboral desarrollada principalmente por "
    "Michael Piore y Peter Doeringer en la década de 1970 explica que las personas no "
    "pueden desplazarse fácilmente entre el sector formal e informal debido a barreras "
    "como el nivel educativo, la experiencia, el acceso a capacitación y la falta de "
    "oportunidades. Por ello, quienes poseen menores ingresos y menor capital humano "
    "tienen una mayor probabilidad de permanecer en empleos informales, mientras que las "
    "personas con mayores niveles de educación, habilidades e ingresos suelen acceder con "
    "más facilidad al sector formal de la economía.")

para(
    "Para Becker (1964), el individuo incurre en gastos de educación al mismo tiempo que "
    "en un costo de oportunidad por permanecer en la población económicamente inactiva y "
    "no recibir renta actual; sin embargo, en el futuro su formación le otorgará la "
    "posibilidad de obtener unos salarios más elevados, pero la productividad de los "
    "empleados depende no sólo de su aptitud y de la inversión que se realiza en ellos, "
    "tanto dentro como fuera del puesto de trabajo, sino también de su motivación y de la "
    "intensidad de su esfuerzo.")

para(
    "Esta teoría muestra que el nivel de educación tiene un fuerte impacto con el nivel de "
    "ingreso futuro de los individuos.")

# ---------------- APORTE DE CLAUDE (rojo) ----------------
para(
    "Asimismo, la investigación se sustenta en tres planteamientos teóricos adicionales "
    "que vinculan los factores socioeconómicos con la informalidad laboral.", color=ROJO)

para(
    "En primer lugar, la **teoría de la elección ocupacional y el enfoque costo-beneficio "
    "de la formalidad** (Maloney, 2004) plantea que la decisión de "
    "insertarse en el sector formal o informal responde a una comparación racional entre "
    "los costos y los beneficios de formalizarse. Los trabajadores con mayores ingresos "
    "—y mayor capacidad de asumir los costos tributarios y administrativos de la "
    "formalidad— tienen una mayor probabilidad de acceder al empleo formal, mientras que "
    "quienes perciben ingresos bajos hallan en la informalidad una estrategia de "
    "subsistencia. Así, un mayor nivel de ingresos contribuye a disminuir la informalidad "
    "laboral.", color=ROJO)

para(
    "En segundo lugar, el **enfoque estructural del desarrollo económico** (Lewis, 1954; "
    "La Porta y Shleifer, 2014) concibe la informalidad como una manifestación del bajo "
    "nivel de ingreso y del limitado desarrollo del sector moderno. Conforme aumentan los "
    "ingresos y se expande el sector formal de alta productividad, la fuerza de trabajo "
    "migra desde las actividades informales de subsistencia hacia el empleo formal. La "
    "evidencia internacional confirma esta relación inversa: la informalidad alcanza cerca "
    "del 88 % de los ocupados en los países de ingreso bajo, frente a solo el 13 % en los "
    "de ingreso alto (Organización Internacional del Trabajo [OIT], 2024). Por ello, a "
    "mayor nivel de ingresos, menor es la informalidad.", color=ROJO)

para(
    "En tercer lugar, el **enfoque del capital humano, la discriminación y la economía de "
    "la familia** (Becker, 1957, 1981; Mincer, 1974) sustenta el efecto de los "
    "determinantes demográficos. La teoría de la discriminación (Becker, 1957) explica "
    "las mayores tasas de informalidad de las mujeres, asociadas a la segregación "
    "ocupacional; la economía de la familia (Becker, 1981) relaciona el estado civil y "
    "las responsabilidades del hogar con la elección de empleos más flexibles e "
    "informales; y la teoría del ciclo de vida laboral (Mincer, 1974) explica el patrón "
    "etario, pues los jóvenes —por su escasa experiencia— y los adultos mayores —por la "
    "depreciación de su capital humano— muestran mayor probabilidad de informalidad. La "
    "evidencia para América Latina identifica el género, la edad y el estado civil como "
    "determinantes significativos de la informalidad (Telles, 1992).", color=ROJO)

# ---------------- REFERENCIAS ----------------
heading("Referencias", size=12)

# Referencias del autor (negro)
referencia(
    "Becker, G. S. (1964). //Human capital: A theoretical and empirical analysis//. "
    "Columbia University Press.", color=NEGRO)
referencia(
    "Doeringer, P. B., & Piore, M. J. (1971). //Internal labor markets and manpower "
    "analysis//. D. C. Heath.", color=NEGRO)
referencia(
    "Hart, K. (1973). Informal income opportunities and urban employment in Ghana. "
    "//The Journal of Modern African Studies//, 11(1), 61-89.", color=NEGRO)
referencia(
    "Tokman, V. E. (1978). An exploration into the nature of informal-formal sector "
    "relationships. //World Development//, 6(9-10), 1065-1075.", color=NEGRO)
referencia(
    "Uribe, J. I., Ortiz, C. H., & Castro, J. A. (2006). //Informalidad laboral en "
    "Colombia, 1988-2000: evolución, teorías y modelos//. Universidad del Valle.",
    color=NEGRO)

# Referencias aportadas por Claude (rojo)
referencia(
    "Becker, G. S. (1957). //The economics of discrimination//. University of Chicago "
    "Press.", color=ROJO)
referencia(
    "Becker, G. S. (1981). //A treatise on the family//. Harvard University Press.",
    color=ROJO)
referencia(
    "La Porta, R., & Shleifer, A. (2014). Informality and development. //Journal of "
    "Economic Perspectives//, 28(3), 109-126. https://doi.org/10.1257/jep.28.3.109",
    color=ROJO)
referencia(
    "Lewis, W. A. (1954). Economic development with unlimited supplies of labour. "
    "//The Manchester School//, 22(2), 139-191.", color=ROJO)
referencia(
    "Maloney, W. F. (2004). Informality revisited. //World Development//, 32(7), "
    "1159-1178. https://doi.org/10.1016/j.worlddev.2004.01.008", color=ROJO)
referencia(
    "Mincer, J. (1974). //Schooling, experience, and earnings//. National Bureau of "
    "Economic Research.", color=ROJO)
referencia(
    "Organización Internacional del Trabajo. (2024). //Perspectivas sociales y del "
    "empleo en el mundo: Tendencias 2024//. OIT.", color=ROJO)
referencia(
    "Telles, E. (1992). Who gets formal sector jobs? Determinants of formal-informal "
    "participation in Brazilian metropolitan areas. //Work and Occupations//, 19(2), "
    "108-127.", color=ROJO)

doc.save(OUT)
print("Guardado:", OUT)
