# -*- coding: utf-8 -*-
"""Genera el Capítulo I (Introducción) en formato Word a partir del markdown fuente.

Replica el formato del Capítulo II ya entregado: portada UNMSM-FCE, encabezado
con el nombre del capítulo, pie de página con autor y número de página,
Times New Roman 12, interlineado 1.5 y texto justificado.
"""
import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

BASE = os.path.dirname(__file__)
SRC = os.path.join(os.path.dirname(BASE), "Capitulo_I_Introduccion.md")
OUT = os.path.join(os.path.dirname(BASE), "CapituloI_Fretel_Anderson.docx")

AUTOR = "Anderson Riquelme Fretel Rojas"
ENCABEZADO = "UNMSM – FCE  |  Capítulo I: Introducción"

doc = Document()

normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(12)
normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

for sec in doc.sections:
    sec.top_margin = Cm(2.5)
    sec.bottom_margin = Cm(2.5)
    sec.left_margin = Cm(3.0)
    sec.right_margin = Cm(2.5)


def _encabezado_pie():
    sec = doc.sections[0]
    head = sec.header.paragraphs[0]
    head.text = ENCABEZADO
    head.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in head.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(9)

    foot = sec.footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = foot.add_run(f"{AUTOR}  |  Página ")
    run.font.name = "Times New Roman"
    run.font.size = Pt(9)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    foot._p.append(fld)


def _runs_con_formato(par, texto, base_bold=False):
    """Convierte **negrita** y *cursiva* del markdown en runs de Word."""
    for trozo in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", texto):
        if not trozo:
            continue
        run = par.add_run()
        if trozo.startswith("**") and trozo.endswith("**"):
            run.text = trozo[2:-2]
            run.bold = True
        elif trozo.startswith("*") and trozo.endswith("*"):
            run.text = trozo[1:-1]
            run.italic = True
        else:
            run.text = trozo
        if base_bold:
            run.bold = True


def parrafo(texto, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=None,
            space_after=6, sangria_lista=False):
    par = doc.add_paragraph()
    par.alignment = align
    par.paragraph_format.space_after = Pt(space_after)
    if sangria_lista:
        par.paragraph_format.left_indent = Cm(0.75)
    _runs_con_formato(par, texto, base_bold=bold)
    if size is not None:
        for run in par.runs:
            run.font.size = Pt(size)
    return par


def titulo(texto, nivel):
    tam = {1: 14, 2: 13, 3: 12}[nivel]
    par = parrafo(texto, bold=True,
                  align=WD_ALIGN_PARAGRAPH.CENTER if nivel == 1 else WD_ALIGN_PARAGRAPH.LEFT,
                  size=tam, space_after=10)
    par.paragraph_format.space_before = Pt(14 if nivel > 1 else 0)
    par.paragraph_format.keep_with_next = True
    return par


def portada():
    for linea, tam, negrita, antes in [
        ("Universidad Nacional Mayor de San Marcos", 16, True, 60),
        ("Facultad de Ciencias Económicas – Escuela Profesional de Economía", 13, False, 6),
        ("Proyecto de Tesis de Grado", 12, False, 30),
        ("CAPÍTULO I", 18, True, 60),
        ("INTRODUCCIÓN", 16, True, 6),
        ('"El efecto de las macrorregiones sobre los ingresos laborales en el Perú: '
         'evidencia de la ENAHO, 2021-2023"', 12, False, 40),
        (f"Autor: {AUTOR}", 12, False, 50),
        ("Docente: Oscar Chavez Polo", 12, False, 6),
        ("Lima, Perú – 2026", 12, False, 40),
    ]:
        par = parrafo(linea, bold=negrita, align=WD_ALIGN_PARAGRAPH.CENTER, size=tam)
        par.paragraph_format.space_before = Pt(antes)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def construir():
    _encabezado_pie()
    portada()

    with open(SRC, encoding="utf-8") as fh:
        lineas = fh.read().splitlines()

    # El bloque de metadatos de la cabecera del markdown (título, autor, docente)
    # ya está representado en la portada; se omite hasta el primer separador.
    i = lineas.index("---") + 1
    en_referencias = False
    for linea in lineas[i:]:
        linea = linea.rstrip()
        if not linea or linea == "---":
            continue
        if linea.startswith("### "):
            titulo(linea[4:], 3)
        elif linea.startswith("## "):
            texto = linea[3:]
            en_referencias = texto.strip().lower() == "referencias"
            titulo(texto, 2)
        elif linea.startswith("# "):
            titulo(linea[2:], 1)
        elif linea.startswith("- "):
            parrafo(linea[2:], sangria_lista=True)
        else:
            par = parrafo(linea)
            if en_referencias:
                # Sangría francesa APA para la lista de referencias.
                par.paragraph_format.left_indent = Cm(1.25)
                par.paragraph_format.first_line_indent = Cm(-1.25)
                par.alignment = WD_ALIGN_PARAGRAPH.LEFT

    doc.save(OUT)
    print(f"Documento generado: {OUT}")


if __name__ == "__main__":
    construir()
